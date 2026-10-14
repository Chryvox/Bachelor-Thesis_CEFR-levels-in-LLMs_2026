"""
cefr_to_excel.py

Reads all CSVs from a folder and creates an Excel file with
  - Info              (documentation of all sheets, columns and metrics)
  - Raw               (all raw data + metadata)
  - Raw_Paired        (both models side by side per text, pivot source; only with exactly 2 models)
  - Mono_Analysis     (formulas on Raw, only if mono data exists)
  - Multi_Analysis    (formulas on Raw, only if multi data exists)
  - Confusion_Matrix  (dropdowns + formulas on Raw)
  - Lists             (hidden, dropdown sources)

Usage:
  python3 cefr_to_excel.py                                -> EVALtoAnalyse/ -> CEFR_Ergebnisse.xlsx
  python3 cefr_to_excel.py FOLDER OUTPUT.xlsx             -> standard folder structure
  python3 cefr_to_excel.py FOLDER OUTPUT.xlsx --generic   -> metadata from the CSV,
                                                             variant = first subfolder
"""
import math
import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------
GENERIC = "--generic" in sys.argv
_args = [a for a in sys.argv[1:] if not a.startswith("--")]
BASE_DIR = Path(_args[0]) if len(_args) > 0 else Path("EVALtoAnalyse")
OUT_FILE = Path(_args[1]) if len(_args) > 1 else Path("CEFR_Ergebnisse.xlsx")

LEVELS = ["A1", "A2", "B1", "B2", "C1", "C2"]
LEVEL_NUM = {l: i + 1 for i, l in enumerate(LEVELS)}   # A1=1 … C2=6
LANG_ORDER = ["de", "en", "ar", "hi"]
TRAIN_ORDER = ["de", "en", "ar", "hi", "en+de", "en+de+ar", "en+de+hi", "en+de+ar+hi"]
SCOPE_ORDER = ["Baseline", "Crosslingual-Mono", "Multi-2", "Multi-3", "Multi-4"]
ALL = "ALL"
CROSS_OPTION = "ALL (except training language)"

# Optional intermediate folder (crosslingual/{train}/{variant}/testOn-…) that counts as a
# regular experiment. Any other variant folder (e.g. half/) is skipped in standard mode.
VARIANT_MAIN = "full"

# Sheet / table names
SH_RAW, SH_PAIRED, SH_CM, SH_LISTS, SH_INFO = "Raw", "Raw_Paired", "Confusion_Matrix", "Lists", "Info"
PAIRED_TABLE = "RawPairedData"

# Fixed column order in Raw. Formulas reference columns by name (COL), not by letter.
RAW_COLS = ["Training_Type", "Training_Scope", "LLM_Model", "Training_Language",
            "Test_Language", "Layer", "Pooling",
            "Level_True", "Level_True_Num", "Level_Predicted", "Level_Predicted_Num",
            "Level_Distance_Abs", "Level_Distance_Signed",
            "Confidence", "Confidence_Level", "Correct", "Is_Cross_Test",
            "Text", "states_filename", "probes_filename", "Source_File"]
COL = {name: get_column_letter(i + 1) for i, name in enumerate(RAW_COLS)}


def rng(name):
    """Whole-column reference on Raw, e.g. Raw!$H:$H"""
    c = COL[name]
    return f"{SH_RAW}!${c}:${c}"


# Styles
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
GROUP_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
GRAY = Font(color="808080", italic=True)
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
THICK_BOTTOM = Border(bottom=Side(style="thick", color="000000"))
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(wrap_text=True, vertical="top")
SIGNED_FMT = "+0.000;-0.000;0.000"   # bias with explicit sign


# Colour scales as functions -> every rule is a separate object
def red_green():
    return ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                          mid_type="num", mid_value=0.5, mid_color="FFEB84",
                          end_type="num", end_value=1, end_color="63BE7B")


def white_blue():
    return ColorScaleRule(start_type="min", start_color="FFFFFF",
                          end_type="max", end_color="5B9BD5")


def blue_white_red():  # for bias (− too low / + too high)
    return ColorScaleRule(start_type="num", start_value=-1, start_color="5B9BD5",
                          mid_type="num", mid_value=0, mid_color="FFFFFF",
                          end_type="num", end_value=1, end_color="F8696B")


# --------------------------------------------------------------------------
# 1) Loading
# --------------------------------------------------------------------------
LAYER_RE = re.compile(r"_l(\d+)\.csv$", re.IGNORECASE)


def parse_path(path: Path) -> dict:
    """Standard mode: metadata from the folder structure."""
    parts = path.relative_to(BASE_DIR).parts
    m = LAYER_RE.search(path.name)
    if not m:
        raise ValueError("no _l{N} in file name")
    layer = int(m.group(1))

    if parts[0] == "baseline" and len(parts) == 4:
        _, test, model, _ = parts
        train, ttype, scope = test, "Mono", "Baseline"
    elif parts[0] == "crosslingual" and len(parts) in (5, 6):
        if len(parts) == 6:   # crosslingual/{train}/{variant}/testOn-{test}/{model}/{file}
            _, train, variant, test_dir, model, _ = parts
            if variant != VARIANT_MAIN:
                raise ValueError(f"variant '{variant}/' is not part of this analysis (evaluated separately)")
        else:
            _, train, test_dir, model, _ = parts
        if not test_dir.startswith("testOn-"):
            raise ValueError(f"unexpected folder '{test_dir}'")
        test = test_dir[len("testOn-"):]
        n = len(train.split("+"))
        ttype = "Mono" if n == 1 else "Multi"
        scope = "Crosslingual-Mono" if n == 1 else f"Multi-{n}"
    else:
        raise ValueError("unexpected folder structure")

    return {
        "Training_Type": ttype,
        "Training_Scope": scope,
        "LLM_Model": model,
        "Training_Language": train,
        "Test_Language": test,
        "Layer": layer,
        "Is_Cross_Test": int(test not in train.split("+")),
        "Source_File": path.relative_to(BASE_DIR).as_posix(),
    }


def short_model(name):
    """CohereLabs/aya-expanse-8b -> aya, Qwen/Qwen3.5-9B -> qwen"""
    s = str(name).lower()
    for alias in ("aya", "qwen"):
        if alias in s:
            return alias
    return str(name).split("/")[-1]


def canon_langs(s):
    """de+en -> en+de (consistent order)"""
    order = {l: i for i, l in enumerate(["en", "de", "ar", "hi"])}
    langs = {l.strip() for l in str(s).split("+") if l.strip()}
    return "+".join(sorted(langs, key=lambda l: (order.get(l, 99), l)))


def parse_generic(path: Path, df: pd.DataFrame) -> dict:
    """Generic mode: metadata from the CSV, variant = first subfolder below BASE_DIR."""
    rel = path.relative_to(BASE_DIR)
    if len(rel.parts) < 2:
        raise ValueError("file must be located in a variant subfolder")
    variant = rel.parts[0]

    vals = {}
    for col in ["LLM_Model", "Training_Language", "Layer"]:
        if col not in df.columns:
            raise ValueError(f"column '{col}' is missing")
        u = df[col].dropna().astype(str).str.strip().unique()
        if len(u) != 1:
            raise ValueError(f"'{col}' is not unique within the file: {list(u)[:5]}")
        vals[col] = u[0]
    if "Test_Language" not in df.columns:
        raise ValueError("column 'Test_Language' is missing")

    train = canon_langs(vals["Training_Language"])
    n = len(train.split("+"))
    return {
        "Training_Type": "Mono" if n == 1 else "Multi",
        "Training_Scope": variant,
        "LLM_Model": short_model(vals["LLM_Model"]),
        "Training_Language": f"{train} ({variant})",
        "Layer": int(float(vals["Layer"])),
        "Source_File": rel.as_posix(),
    }


def load_all() -> pd.DataFrame:
    frames = []
    for path in sorted(BASE_DIR.rglob("*.csv")):
        df = pd.read_csv(path, encoding="utf-8-sig")
        try:
            meta = parse_generic(path, df) if GENERIC else parse_path(path)
        except ValueError as e:
            print(f"[SKIP] {path}: {e}")
            continue

        if GENERIC:
            # Test language per row from the CSV, cross test against the pure training languages
            train_set = set(meta["Training_Language"].split(" (")[0].split("+"))
            df["Test_Language"] = df["Test_Language"].astype(str).str.strip()
            df["Is_Cross_Test"] = (~df["Test_Language"].isin(train_set)).astype(int)
        else:
            # The path is the "truth" – only report REAL differences to the CSV
            # (full HF model name or a different language order are not a problem)
            for col in ["LLM_Model", "Training_Language", "Test_Language", "Layer"]:
                if col not in df.columns:
                    continue
                s = df[col].astype(str).str.strip()
                if col == "Layer":
                    diff = pd.to_numeric(df[col], errors="coerce") != meta[col]
                elif col == "LLM_Model":
                    diff = s.map(short_model) != meta[col]
                elif col == "Training_Language":
                    diff = s.map(canon_langs) != canon_langs(meta[col])
                else:
                    diff = s != str(meta[col])
                if diff.any():
                    print(f"[WARN] {path.name}: '{col}' in CSV = {df.loc[diff, col].iloc[0]!r}, "
                          f"path = {meta[col]!r} -> using path value")

        for k, v in meta.items():
            df[k] = v
        frames.append(df)
        print(f"[OK]   {meta['Source_File']}  ({len(df)} rows)")

    if not frames:
        sys.exit(f"No (valid) CSVs found below {BASE_DIR.resolve()}.")

    raw = pd.concat(frames, ignore_index=True)
    for c in RAW_COLS:
        if c not in raw.columns:
            raw[c] = None

    # Normalise levels
    for c in ["Level_True", "Level_Predicted"]:
        raw[c] = raw[c].astype(str).str.strip().str.upper()
    bad = ~raw["Level_True"].isin(LEVELS) | ~raw["Level_Predicted"].isin(LEVELS)
    if bad.any():
        print(f"[WARN] {bad.sum()} rows with unknown level are dropped")
        raw = raw[~bad].copy()

    # Recompute derived level columns consistently (numeric)
    true_num = raw["Level_True"].map(LEVEL_NUM)
    pred_num = raw["Level_Predicted"].map(LEVEL_NUM)
    signed = pred_num - true_num                     # + = too high, - = too low
    derived = {
        "Level_True_Num": true_num,
        "Level_Predicted_Num": pred_num,
        "Level_Distance_Signed": signed,
        "Level_Distance_Abs": signed.abs(),
        "Correct": (signed == 0).astype(int),
    }
    for c, new in derived.items():
        orig = raw[c]
        if c == "Correct":  # in case the CSV contains "True"/"False"
            orig = orig.astype(str).str.strip().str.lower().map(
                {"true": 1, "1": 1, "1.0": 1, "false": 0, "0": 0, "0.0": 0})
        else:
            orig = pd.to_numeric(orig, errors="coerce")
        n_diff = int(((orig != new) & orig.notna()).sum())
        if n_diff:
            print(f"[WARN] {n_diff} rows: '{c}' in CSV did not match the levels -> recomputed")
        raw[c] = new.astype(int)

    raw["Confidence"] = pd.to_numeric(raw["Confidence"], errors="coerce")

    # Take Confidence_Level from the CSV, only compute missing values
    conf_lvl = pd.to_numeric(raw["Confidence_Level"], errors="coerce")
    missing = conf_lvl.isna() & raw["Confidence"].notna()
    if missing.any():
        print(f"[INFO] {int(missing.sum())} rows without Confidence_Level -> computed from Confidence")
        conf_lvl[missing] = (raw.loc[missing, "Confidence"] * 10).round()
    raw["Confidence_Level"] = conf_lvl

    raw["Pooling"] = raw["Pooling"].fillna("n/a").astype(str).str.strip()
    raw["Layer"] = raw["Layer"].astype(int)

    if len(raw) + 1 > 1_048_576:
        sys.exit(f"{len(raw)} rows – too many for one Excel sheet.")
    print(f"\nTotal: {len(raw)} rows from {len(frames)} files.\n")
    return raw[RAW_COLS].reset_index(drop=True)


# --------------------------------------------------------------------------
# 2) Raw sheet
# --------------------------------------------------------------------------
def clean(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)[:32767]
    if hasattr(v, "item"):  # numpy -> python
        v = v.item()
        if isinstance(v, float) and pd.isna(v):
            return None
    return v


def append_rows(ws, rows):
    """Write rows; texts starting with '=' are stored as text, NOT as formulas."""
    for row in rows:
        vals = [clean(v) for v in row]
        ws.append(vals)
        for j, v in enumerate(vals):
            if isinstance(v, str) and v.startswith("="):
                ws.cell(ws.max_row, j + 1).data_type = "s"


def write_raw(ws, raw):
    ws.append(RAW_COLS)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    append_rows(ws, raw.itertuples(index=False, name=None))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"Training_Scope": 18, "Training_Language": 18, "Level_True_Num": 15,
              "Level_Predicted_Num": 19, "Level_Distance_Abs": 18,
              "Level_Distance_Signed": 21, "Confidence_Level": 16, "Text": 60,
              "states_filename": 30, "probes_filename": 30, "Source_File": 50}
    for name, letter in COL.items():
        ws.column_dimensions[letter].width = widths.get(name, 13)


# --------------------------------------------------------------------------
# 3) Raw_Paired sheet: both models per text in one row (pivot source)
# --------------------------------------------------------------------------
# Key columns are identical for both models; value columns exist once per model.
PAIR_KEY = ["Training_Type", "Training_Scope", "Training_Language", "Test_Language",
            "Is_Cross_Test", "Level_True", "Level_True_Num", "Text"]
PAIR_VALS = ["Level_Predicted", "Level_Predicted_Num", "Confidence", "Confidence_Level",
             "Level_Distance_Abs", "Level_Distance_Signed", "Correct", "Layer", "Source_File"]


def model_labels(raw):
    """Sorted model names + capitalised prefixes, e.g. (aya, qwen, Aya, Qwen)."""
    m1, m2 = sorted(raw["LLM_Model"].unique())
    return m1, m2, m1.capitalize(), m2.capitalize()


def build_paired(ws, raw):
    m1, m2, P1, P2 = model_labels(raw)

    parts = {}
    for m in (m1, m2):
        d = raw.loc[raw["LLM_Model"] == m, PAIR_KEY + PAIR_VALS].copy()
        # Pair genuine duplicate texts (e.g. enToDe) 1:1 instead of multiplying them 2×2
        d["_occ"] = d.groupby(PAIR_KEY).cumcount()
        parts[m] = d.rename(columns={c: f"{m.capitalize()}_{c}" for c in PAIR_VALS})

    pair = parts[m1].merge(parts[m2], on=PAIR_KEY + ["_occ"], how="inner")
    for m in (m1, m2):
        n_un = len(parts[m]) - len(pair)
        if n_un:
            print(f"[WARN] {SH_PAIRED}: {n_un} rows of '{m}' have no partner in the other model (left out)")
    print(f"[OK]   {SH_PAIRED}: {len(pair)} paired texts ({P1} × {P2})")

    if pair.empty:
        ws["A1"] = "No matching texts between the two models."
        return

    # Derived columns
    pair["Distance_Pair"] = (pair[f"{P1}_Level_Distance_Abs"].astype(str) + "/"
                             + pair[f"{P2}_Level_Distance_Abs"].astype(str))
    c1, c2 = pair[f"{P1}_Correct"] == 1, pair[f"{P2}_Correct"] == 1
    pair["Correct_Pair"] = "both wrong"
    pair.loc[c1 & ~c2, "Correct_Pair"] = f"only {P1} correct"
    pair.loc[~c1 & c2, "Correct_Pair"] = f"only {P2} correct"
    pair.loc[c1 & c2, "Correct_Pair"] = "both correct"
    pair["Same_Prediction"] = (pair[f"{P1}_Level_Predicted"] == pair[f"{P2}_Level_Predicted"]).astype(int)

    # Same sort order as in the analysis sheets
    pair = (pair.assign(
                _s=pair["Training_Scope"].map({v: i for i, v in enumerate(SCOPE_ORDER)}).fillna(99),
                _t=pair["Training_Language"].map({v: i for i, v in enumerate(TRAIN_ORDER)}).fillna(99),
                _e=pair["Test_Language"].map({v: i for i, v in enumerate(LANG_ORDER)}).fillna(99))
            .sort_values(["_s", "Training_Scope", "_t", "Training_Language", "_e", "Test_Language"]))

    both = lambda c: [f"{P1}_{c}", f"{P2}_{c}"]
    cols = (["Training_Type", "Training_Scope", "Training_Language", "Test_Language",
             "Is_Cross_Test", "Level_True", "Level_True_Num"]
            + both("Level_Predicted") + both("Level_Predicted_Num")
            + both("Confidence_Level") + both("Confidence")
            + both("Level_Distance_Abs") + both("Level_Distance_Signed") + both("Correct")
            + ["Distance_Pair", "Correct_Pair", "Same_Prediction", "Text"]
            + both("Layer") + both("Source_File"))

    ws.append(cols)
    append_rows(ws, pair[cols].itertuples(index=False, name=None))

    # Real Excel table -> ideal as pivot source
    last = get_column_letter(len(cols))
    tab = Table(displayName=PAIRED_TABLE, ref=f"A1:{last}{len(pair) + 1}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 60 if c == "Text" else max(12, min(len(c) + 2, 28))


# --------------------------------------------------------------------------
# 4) Lists sheet (dropdown sources)
# --------------------------------------------------------------------------
def ordered(values, order):
    vals = list(dict.fromkeys(values))
    return [v for v in order if v in vals] + sorted(v for v in vals if v not in order)


def write_lists(ws, raw):
    lists = {
        "model": ("Model", [ALL] + sorted(raw["LLM_Model"].unique())),
        "train": ("Training Language", [ALL] + ordered(raw["Training_Language"].unique(), TRAIN_ORDER)),
        "test": ("Test Language", ordered(raw["Test_Language"].unique(), LANG_ORDER) + [ALL, CROSS_OPTION]),
    }
    refs = {}
    for ci, (key, (title, vals)) in enumerate(lists.items(), start=1):
        letter = get_column_letter(ci)
        ws.cell(1, ci, title).font = BOLD
        for ri, v in enumerate(vals, start=2):
            ws.cell(ri, ci, v)
        refs[key] = f"={SH_LISTS}!${letter}$2:${letter}${len(vals) + 1}"
        ws.column_dimensions[letter].width = 32
    ws.sheet_state = "hidden"
    return refs


# --------------------------------------------------------------------------
# 5) Mono_Analysis / Multi_Analysis
# --------------------------------------------------------------------------
KEY_COLS = ["Training_Scope", "LLM_Model", "Training_Language", "Test_Language", "Layer", "Pooling"]


def analysis_spec():
    spec = [("Overview", h) for h in
            ["Training_Scope", "Model", "Training Language(s)", "Test Language", "Layer", "Pooling",
             "Test Samples"]]
    spec += [("Metrics", h) for h in
             ["Correct", "Accuracy", "Accuracy X/N", "Adjacent Correct", "Adjacent Acc (±1)",
              "Adjacent X/N", "MAE (Level)",
              "Bias (Signed)", "% Overestimated", "% Underestimated",
              "Conf Avg", "Conf Min", "Conf Max"]]
    for group, prefix in [("Precision", "P"), ("Recall", "R"), ("F1-Score", "F1")]:
        spec += [(group, f"{prefix} {l}") for l in LEVELS]
    spec += [("Macro Avg", h) for h in ["Macro P", "Macro R", "Macro F1"]]
    for group in ["Support", "TP", "FP", "FN", "TN"]:
        spec += [(group, f"{group} {l}") for l in LEVELS]
    return spec


def build_analysis(ws, raw, ttype):
    spec = analysis_spec()
    L = {h: get_column_letter(i + 1) for i, (_, h) in enumerate(spec)}
    last_col = get_column_letter(len(spec))

    ws["A1"] = f"{ttype} Analysis – all values are formulas on sheet 'Raw' (explanations: sheet 'Info')"
    ws["A1"].font = Font(bold=True, size=14)

    # Row 2: groups (merged), row 3: column headers
    start = 1
    for i in range(1, len(spec) + 1):
        if i == len(spec) or spec[i][0] != spec[start - 1][0]:
            ws.cell(2, start, spec[start - 1][0])
            if i > start:
                ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=i)
            start = i + 1
    for i, (_, h) in enumerate(spec, start=1):
        g = ws.cell(2, i)
        g.fill, g.font, g.alignment = GROUP_FILL, BOLD, CENTER
        c = ws.cell(3, i, h)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, CENTER

    keys = raw.loc[raw["Training_Type"] == ttype, KEY_COLS].drop_duplicates()
    if keys.empty:
        ws["A4"] = f"No {ttype} data in Raw."
        return
    keys = (keys.assign(
                _s=keys["Training_Scope"].map({v: i for i, v in enumerate(SCOPE_ORDER)}).fillna(99),
                _t=keys["Training_Language"].map({v: i for i, v in enumerate(TRAIN_ORDER)}).fillna(99),
                _e=keys["Test_Language"].map({v: i for i, v in enumerate(LANG_ORDER)}).fillna(99))
            # Scope -> training language -> test language -> then model;
            # unknown scopes / training languages are sorted alphabetically
            .sort_values(["_s", "Training_Scope", "_t", "Training_Language", "_e",
                          "LLM_Model", "Layer", "Pooling"]))

    pct = {"Accuracy", "Adjacent Acc (±1)", "% Overestimated", "% Underestimated"}
    dec = {"MAE (Level)", "Conf Avg", "Conf Min", "Conf Max", "Macro P", "Macro R", "Macro F1"} \
        | {f"{p} {l}" for p in ("P", "R", "F1") for l in LEVELS}

    r = 3
    prev_grp, prev_test = None, None
    for k in keys.itertuples(index=False):
        scope, model, train, test, layer, pooling = k[:6]
        grp = (scope, train)
        if prev_grp is not None and grp != prev_grp:
            r += 1  # blank row between training languages
        elif prev_grp is not None and test != prev_test:
            # thick line below the last block of the same training language
            for c in range(1, len(spec) + 1):
                ws.cell(r, c).border = THICK_BOTTOM
        r += 1
        prev_grp, prev_test = grp, test

        crit = (f"{rng('Training_Scope')},$A{r},{rng('LLM_Model')},$B{r},"
                f"{rng('Training_Language')},$C{r},{rng('Test_Language')},$D{r},"
                f"{rng('Layer')},$E{r},{rng('Pooling')},$F{r}")
        X = lambda h: f"{L[h]}{r}"
        N = X("Test Samples")

        f = {
            "Training_Scope": scope, "Model": model, "Training Language(s)": train,
            "Test Language": test, "Layer": int(layer), "Pooling": pooling,
            "Test Samples": f"=COUNTIFS({crit})",
            "Correct": f"=COUNTIFS({crit},{rng('Correct')},1)",
            "Accuracy": f'=IF({N}=0,"",{X("Correct")}/{N})',
            "Accuracy X/N": f'={X("Correct")}&"/"&{N}',
            "Adjacent Correct": f'=COUNTIFS({crit},{rng("Level_Distance_Abs")},"<=1")',
            "Adjacent Acc (±1)": f'=IF({N}=0,"",{X("Adjacent Correct")}/{N})',
            "Adjacent X/N": f'={X("Adjacent Correct")}&"/"&{N}',
            "MAE (Level)": f'=IFERROR(AVERAGEIFS({rng("Level_Distance_Abs")},{crit}),"")',
            "Bias (Signed)": f'=IFERROR(AVERAGEIFS({rng("Level_Distance_Signed")},{crit}),"")',
            "% Overestimated": f'=IF({N}=0,"",COUNTIFS({crit},{rng("Level_Distance_Signed")},">0")/{N})',
            "% Underestimated": f'=IF({N}=0,"",COUNTIFS({crit},{rng("Level_Distance_Signed")},"<0")/{N})',
            "Conf Avg": f'=IFERROR(AVERAGEIFS({rng("Confidence")},{crit}),"")',
            # MINIFS/MAXIFS need the _xlfn. prefix, otherwise #NAME?
            "Conf Min": f'=IF({N}=0,"",_xlfn.MINIFS({rng("Confidence")},{crit}))',
            "Conf Max": f'=IF({N}=0,"",_xlfn.MAXIFS({rng("Confidence")},{crit}))',
        }
        for l in LEVELS:
            tp, sup, fp = X(f"TP {l}"), X(f"Support {l}"), X(f"FP {l}")
            p, rc = X(f"P {l}"), X(f"R {l}")
            f[f"Support {l}"] = f'=COUNTIFS({crit},{rng("Level_True")},"{l}")'
            f[f"TP {l}"] = f'=COUNTIFS({crit},{rng("Level_True")},"{l}",{rng("Level_Predicted")},"{l}")'
            f[f"FP {l}"] = f'=COUNTIFS({crit},{rng("Level_Predicted")},"{l}")-{tp}'
            f[f"FN {l}"] = f"={sup}-{tp}"
            f[f"TN {l}"] = f"={N}-{tp}-{fp}-{X(f'FN {l}')}"
            f[f"P {l}"] = f"=IF({tp}+{fp}=0,0,{tp}/({tp}+{fp}))"
            f[f"R {l}"] = f"=IF({sup}=0,0,{tp}/{sup})"
            f[f"F1 {l}"] = f"=IF({p}+{rc}=0,0,2*{p}*{rc}/({p}+{rc}))"

        # Macro avg like sklearn: only classes that occur in y_true OR y_pred
        sup_r = f"{L['Support A1']}{r}:{L['Support C2']}{r}"
        fp_r = f"{L['FP A1']}{r}:{L['FP C2']}{r}"
        present = f"SUMPRODUCT(--(({sup_r})+({fp_r})>0))"
        for name, pre in [("Macro P", "P"), ("Macro R", "R"), ("Macro F1", "F1")]:
            f[name] = f'=IFERROR(SUM({L[pre + " A1"]}{r}:{L[pre + " C2"]}{r})/{present},"")'

        for h, v in f.items():
            cell = ws[X(h)]
            cell.value = v
            if h in pct:
                cell.number_format = "0.0%"
            elif h in dec:
                cell.number_format = "0.000"
            elif h == "Bias (Signed)":
                cell.number_format = SIGNED_FMT

    last_row = r
    for h in ["Accuracy", "Adjacent Acc (±1)", "Macro F1"] + [f"F1 {l}" for l in LEVELS]:
        ws.conditional_formatting.add(f"{L[h]}4:{L[h]}{last_row}", red_green())
    ws.conditional_formatting.add(f"{L['Bias (Signed)']}4:{L['Bias (Signed)']}{last_row}",
                                  blue_white_red())

    ws.auto_filter.ref = f"A3:{last_col}{last_row}"
    ws.freeze_panes = "G4"
    for i, (_, h) in enumerate(spec, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(9, min(len(h) + 3, 20))


# --------------------------------------------------------------------------
# 6) Confusion_Matrix
# --------------------------------------------------------------------------
def build_confusion(ws, refs):
    ws["A1"] = "Confusion Matrix – live from 'Raw'"
    ws["A1"].font = Font(bold=True, size=14)

    # Dropdowns
    for row, (label, key) in enumerate([("Model:", "model"),
                                        ("Training Language(s):", "train"),
                                        ("Test Language:", "test")], start=3):
        ws.cell(row, 1, label).font = BOLD
        cell = ws.cell(row, 2, ALL)
        cell.fill, cell.border, cell.font = INPUT_FILL, BOX, BOLD
        dv = DataValidation(type="list", formula1=refs[key], allow_blank=False)
        dv.errorTitle, dv.error = "Invalid value", "Please select a value from the list."
        ws.add_data_validation(dv)
        dv.add(cell)

    # Internal criteria (translate dropdown values into COUNTIFS criteria)
    ws["D2"] = "Filter criteria (internal, do not edit)"
    ws["D2"].font = GRAY
    helpers = [
        ("Model", f'=IF($B$3="{ALL}","*",$B$3)'),
        ("Training", f'=IF($B$4="{ALL}","*",$B$4)'),
        ("Test", f'=IF(OR($B$5="{ALL}",$B$5="{CROSS_OPTION}"),"*",$B$5)'),
        ("Is_Cross", f'=IF($B$5="{CROSS_OPTION}",1,">=0")'),
    ]
    for i, (lab, frm) in enumerate(helpers, start=3):
        ws.cell(i, 4, lab).font = GRAY
        ws.cell(i, 5, frm).font = GRAY

    ws["A7"] = "Note: layers & pooling are aggregated – each layer file counts separately."
    ws["A7"].font = GRAY

    crit = (f"{rng('LLM_Model')},$E$3,{rng('Training_Language')},$E$4,"
            f"{rng('Test_Language')},$E$5,{rng('Is_Cross_Test')},$E$6")

    # Absolute matrix: rows 9–14 (true), columns B–G (predicted)
    top, first = 8, 9
    ws.cell(top, 1, "True ↓ / Pred →")
    ws.cell(top, 8, "Σ True")
    for j, l in enumerate(LEVELS):
        ws.cell(top, 2 + j, l)
        ws.cell(first + j, 1, l)
    for c in range(1, 9):
        ws.cell(top, c).fill, ws.cell(top, c).font = HEAD_FILL, HEAD_FONT
        ws.cell(top, c).alignment = CENTER

    for i in range(6):
        r = first + i
        ws.cell(r, 1).font = BOLD
        ws.cell(r, 1).fill = GROUP_FILL
        for j in range(6):
            col = get_column_letter(2 + j)
            c = ws.cell(r, 2 + j,
                        f"=COUNTIFS({crit},{rng('Level_True')},$A{r},{rng('Level_Predicted')},{col}${top})")
            c.border, c.alignment = BOX, CENTER
            if i == j:
                c.font = BOLD
        ws.cell(r, 8, f"=SUM(B{r}:G{r})").font = BOLD

    sum_row = first + 6
    ws.cell(sum_row, 1, "Σ Pred").font = BOLD
    for j in range(7):
        col = get_column_letter(2 + j)
        ws.cell(sum_row, 2 + j, f"=SUM({col}{first}:{col}{first + 5})").font = BOLD
    ws.conditional_formatting.add(f"B{first}:G{first + 5}", white_blue())

    # Row-normalised matrix (K–Q)
    ws.cell(top, 11, "Row %")
    for j, l in enumerate(LEVELS):
        ws.cell(top, 12 + j, l)
    for c in range(11, 18):
        ws.cell(top, c).fill, ws.cell(top, c).font = HEAD_FILL, HEAD_FONT
        ws.cell(top, c).alignment = CENTER
    for i in range(6):
        r = first + i
        ws.cell(r, 11, LEVELS[i]).font = BOLD
        ws.cell(r, 11).fill = GROUP_FILL
        for j in range(6):
            src = get_column_letter(2 + j)
            c = ws.cell(r, 12 + j, f'=IF($H{r}=0,"",{src}{r}/$H{r})')
            c.number_format, c.border, c.alignment = "0.0%", BOX, CENTER
    ws.conditional_formatting.add(f"L{first}:Q{first + 5}", white_blue())

    # Classification report
    rep = sum_row + 3
    for c, h in enumerate(["Class", "Precision", "Recall", "F1", "Support", "Pred Count"], start=1):
        cell = ws.cell(rep, c, h)
        cell.fill, cell.font, cell.alignment = HEAD_FILL, HEAD_FONT, CENTER
    for i, l in enumerate(LEVELS):
        r, m = rep + 1 + i, first + i
        d = get_column_letter(2 + i)
        ws.cell(r, 1, l).font = BOLD
        ws.cell(r, 2, f"=IF({d}${sum_row}=0,0,{d}{m}/{d}${sum_row})")
        ws.cell(r, 3, f"=IF($H{m}=0,0,{d}{m}/$H{m})")
        ws.cell(r, 4, f"=IF(B{r}+C{r}=0,0,2*B{r}*C{r}/(B{r}+C{r}))")
        ws.cell(r, 5, f"=$H{m}")
        ws.cell(r, 6, f"={d}${sum_row}")
        for c in (2, 3, 4):
            ws.cell(r, c).number_format = "0.000"
    a, b = rep + 1, rep + 6
    mr = b + 1
    present = f"SUMPRODUCT(--((E{a}:E{b})+(F{a}:F{b})>0))"
    ws.cell(mr, 1, "Macro Avg").font = BOLD
    for c, col in [(2, "B"), (3, "C"), (4, "D")]:
        ws.cell(mr, c, f'=IFERROR(SUM({col}{a}:{col}{b})/{present},"")').number_format = "0.000"
        ws.cell(mr, c).font = BOLD
    ws.cell(mr, 5, f"=SUM(E{a}:E{b})").font = BOLD
    ws.conditional_formatting.add(f"D{a}:D{b}", red_green())

    # Overall metrics (row i = true, column j = predicted)
    cell_ref = lambda i, j: f"{get_column_letter(2 + j)}{first + i}"
    pairs = [(i, j) for i in range(6) for j in range(6)]
    diag = "+".join(cell_ref(i, i) for i in range(6))
    adj = "+".join(cell_ref(i, j) for i, j in pairs if abs(i - j) <= 1)
    mae = "+".join(f"{abs(i - j)}*{cell_ref(i, j)}" for i, j in pairs if i != j)
    # Bias & over-/underestimation (above diagonal = too high, below = too low)
    over_w = "+".join(f"{j - i}*{cell_ref(i, j)}" for i, j in pairs if j > i)
    under_w = "+".join(f"{i - j}*{cell_ref(i, j)}" for i, j in pairs if j < i)
    over = "+".join(cell_ref(i, j) for i, j in pairs if j > i)
    under = "+".join(cell_ref(i, j) for i, j in pairs if j < i)

    g = mr + 2
    total = f"H{sum_row}"
    for k, (lab, frm, fmt) in enumerate([
        ("N (Predictions)", f"={total}", "0"),
        ("Accuracy (exact)", f'=IF({total}=0,"",({diag})/{total})', "0.0%"),
        ("Adjacent Acc (±1)", f'=IF({total}=0,"",({adj})/{total})', "0.0%"),
        ("MAE (Level)", f'=IF({total}=0,"",({mae})/{total})', "0.000"),
        ("Bias (Signed)", f'=IF({total}=0,"",(({over_w})-({under_w}))/{total})', SIGNED_FMT),
        ("% Overestimated", f'=IF({total}=0,"",({over})/{total})', "0.0%"),
        ("% Underestimated", f'=IF({total}=0,"",({under})/{total})', "0.0%"),
    ]):
        ws.cell(g + k, 1, lab).font = BOLD
        ws.cell(g + k, 2, frm).number_format = fmt
    ws.conditional_formatting.add(f"B{g + 4}", blue_white_red())

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    for c in "CFGHKLMNOPQ":
        ws.column_dimensions[c].width = max(ws.column_dimensions[c].width or 0, 10)


# --------------------------------------------------------------------------
# 7) Info sheet
# --------------------------------------------------------------------------
RAW_INFO = {
    "Training_Type": ("Path", "Mono = trained on exactly one language (incl. baseline), "
                              "Multi = trained on several languages."),
    "Training_Scope": ("Path", "Baseline = folder baseline/ (training language = test language). "
                               "Crosslingual-Mono = one training language, tested on other languages. "
                               "Multi-2 / Multi-3 / Multi-4 = number of training languages. "
                               "In generic mode: name of the variant subfolder."),
    "LLM_Model": ("Path", "LLM that produced the hidden states. Short name (e.g. aya, qwen). "
                          "The full name from the CSV (e.g. Qwen/Qwen3.5-9B) is replaced so that filters "
                          "and dropdowns are consistent."),
    "Training_Language": ("Path", "Language(s) the probe was trained on, e.g. de or en+de+ar. "
                                  "Consistent notation (en+de instead of de+en). In generic mode the variant "
                                  "is appended, e.g. en+de (half)."),
    "Test_Language": ("Path", "Language of the tested text (from baseline/{x} or testOn-{x}; "
                              "in generic mode per row from the CSV)."),
    "Layer": ("File name", "Layer the states were taken from (_l16 → 16). In generic mode from the CSV."),
    "Pooling": ("CSV", "Pooling method over the tokens (e.g. mean)."),
    "Level_True": ("CSV", "True CEFR label (A1–C2)."),
    "Level_True_Num": ("Recomputed", "True label as a number: A1=1, A2=2, B1=3, B2=4, C1=5, C2=6."),
    "Level_Predicted": ("CSV", "Predicted CEFR label (A1–C2)."),
    "Level_Predicted_Num": ("Recomputed", "Prediction as a number (A1=1 … C2=6)."),
    "Level_Distance_Abs": ("Recomputed", "|Predicted_Num − True_Num|, range 0–5. 0 = correct, "
                                         "B1↔B2 = 1, A1↔C2 = 5. Basis for MAE and adjacent accuracy."),
    "Level_Distance_Signed": ("Recomputed", "Predicted_Num − True_Num, range −5 … +5. "
                                            "+ = rated too high, − = rated too low, 0 = correct. Basis for bias."),
    "Confidence": ("CSV", "Confidence of the prediction (0–1)."),
    "Confidence_Level": ("CSV (missing → computed)", "round(Confidence × 10) → bucket 0–10. "
                                                     "For your own pivot analyses (e.g. accuracy per bucket)."),
    "Correct": ("Recomputed", "1 = correct (distance 0), 0 = wrong. Mean = accuracy."),
    "Is_Cross_Test": ("Path", "1 = test language was NOT among the training languages (true cross-lingual "
                              "test). 0 = test language was part of training (incl. baseline)."),
    "Text": ("CSV", "The original classified text."),
    "states_filename": ("CSV", "Name of the test states file used."),
    "probes_filename": ("CSV", "Name of the probe file used."),
    "Source_File": ("Script", "Relative path of the CSV (below the input folder) the row comes from."),
}

SHEETS_INFO = [
    ("Info", "This page", "Explanation of all sheets, columns and metrics."),
    ("Raw", "All predictions (values)",
     "One row = one prediction for one text in one experiment. Only data source for all formulas. "
     "Do not change the column order."),
    (SH_PAIRED, "Both models side by side (values)",
     "One row = one text in one experiment, with the predictions of both models next to each other "
     "(matched on scope, training language, test language, true level and text). Stored as the Excel "
     f"table '{PAIRED_TABLE}' – intended as a pivot source. Static values: rerun the script after "
     "changes to Raw. Only created if there are exactly two models."),
    ("Mono_Analysis", "Metrics per experiment (mono)",
     "Training_Type = Mono. One row = one experiment (scope + model + training language + test language + "
     "layer + pooling). All metrics are formulas on Raw. Omitted if there is no mono data."),
    ("Multi_Analysis", "Metrics per experiment (multi)",
     "Training_Type = Multi. Same layout as Mono_Analysis. Omitted if there is no multi data."),
    ("Confusion_Matrix", "Interactive 6×6 matrix",
     "Dropdowns for model, training language(s) and test language. Matrix, classification report and "
     "overall metrics are computed live from Raw."),
    ("Lists (hidden)", "Dropdown sources",
     "Generated from the available data when the script runs. To show it: right-click a sheet tab → Unhide."),
]


def paired_info(P1, P2):
    """Column descriptions for Raw_Paired (model prefixes depend on the data)."""
    pp = lambda c: f"{P1}_{c} / {P2}_{c}"
    return [
        ("Training_Type … Is_Cross_Test", "Experiment key, identical to Raw."),
        ("Level_True / Level_True_Num", "True CEFR label of the text as text and as a number "
                                        "(A1=1 … C2=6), identical for both models."),
        (pp("Level_Predicted"), "Predicted level of each model (text)."),
        (pp("Level_Predicted_Num"), "Predicted level of each model as a number (A1=1 … C2=6)."),
        (pp("Confidence_Level"), "Confidence bucket 0–10 of each model. Typical pivot rows: "
                                 f"{P1}_Confidence_Level, then {P2}_Confidence_Level."),
        (pp("Confidence"), "Raw confidence (0–1) of each model."),
        (pp("Level_Distance_Abs"), "Absolute level distance of each model (0–5)."),
        (pp("Level_Distance_Signed"), "Signed level distance of each model (+ too high, − too low)."),
        (pp("Correct"), "1 = correct, 0 = wrong, per model."),
        ("Distance_Pair", f"Combination of both absolute distances as text '{P1}/{P2}', "
                          f"e.g. 0/1 = {P1} correct, {P2} one level off."),
        ("Correct_Pair", f"'both correct', 'only {P1} correct', 'only {P2} correct' or 'both wrong'."),
        ("Same_Prediction", "1 = both models predicted the same level (right or wrong), 0 = different."),
        ("Text", "The classified text."),
        (pp("Layer") + ", " + pp("Source_File"), "Layer and source CSV of each model."),
        ("Pivot example", f"Insert → PivotTable with source '{PAIRED_TABLE}'. Filters: Training_Type, "
                          f"Training_Language, Test_Language. Rows: {P1}_Confidence_Level, "
                          f"{P2}_Confidence_Level. Columns: Distance_Pair or Correct_Pair (or "
                          f"{P1}_/{P2}_Level_Distance_Abs). Values: count of Text. Always filter to one "
                          "training and test language, otherwise several experiments are added up."),
        ("Unmatched texts", "Texts that exist for only one model in an experiment are left out "
                            "(reported as [WARN] in the log)."),
    ]


METRIC_INFO = [
    ("Test Samples (N)", "Number of Raw rows of the experiment", "Population for all other metrics."),
    ("Correct / Accuracy", "Count of Correct = 1  /  N", "Share of texts classified exactly right."),
    ("Adjacent Acc (±1)", "Count of Level_Distance_Abs ≤ 1  /  N",
     "Share of predictions at most one level off. Accounts for the ordinal nature of CEFR levels."),
    ("… X/N", "Text, e.g. 187/246", "Shows the underlying absolute numbers."),
    ("MAE (Level)", "Mean of Level_Distance_Abs",
     "Average deviation in levels. 0 = perfect, 1 = on average one level off."),
    ("Bias (Signed)", "Mean of Level_Distance_Signed",
     "Systematic tendency: > 0 = rates too high on average, < 0 = too low, ≈ 0 = no tendency. "
     "Over- and underestimations cancel out in the bias, so always read it together with MAE."),
    ("% Overestimated / % Underestimated", "Share of Signed > 0 or < 0",
     "Share of texts rated too high or too low. Together with accuracy = 100 %."),
    ("Conf Avg / Min / Max", "Mean / minimum / maximum of Confidence",
     "How confident the probe is on average and in the extremes."),
    ("Precision (class i)", "TPᵢ / (TPᵢ + FPᵢ)",
     "When the model predicts i: how often is that right? 0 if i was never predicted."),
    ("Recall (class i)", "TPᵢ / Supportᵢ", "Which share of the true i texts is recognised?"),
    ("F1-Score (class i)", "2 · P · R / (P + R)", "Harmonic mean of precision and recall."),
    ("Macro P / R / F1", "Mean over all classes with Support + FP > 0 (as in sklearn)",
     "Every class counts equally regardless of frequency, so rare classes carry more weight."),
    ("Support (class i)", "Count of Level_True = i", "Number of true i texts in the experiment."),
    ("TP (class i)", "True = i and Pred = i", "Correctly recognised as i."),
    ("FP (class i)", "Pred = i, but True ≠ i", "Wrongly classified as i."),
    ("FN (class i)", "True = i, but Pred ≠ i", "True i text not recognised as i."),
    ("TN (class i)", "True ≠ i and Pred ≠ i", "Correctly treated as not-i (one-vs-rest)."),
]

CM_INFO = [
    ("Dropdowns (B3–B5)", f"Model, training language(s), test language. '{ALL}' = no filter on this dimension."),
    (CROSS_OPTION, "Filters on Is_Cross_Test = 1, i.e. only tests on languages that were not part of "
                   "training. Works automatically for every training language."),
    ("Helper cells (E3–E6, grey)", "Translate the dropdown selection into COUNTIFS criteria. Do not edit."),
    ("Absolute matrix (B9:G14)", "Rows = true level, columns = predicted level, cell = number of predictions. "
                                 "Diagonal = correct. Above the diagonal = rated too high, below = rated too "
                                 "low. Σ True = support, Σ Pred = number of predictions per class."),
    ("Row % (L9:Q14)", "Each row normalised to 100 %: how are the true texts of a level distributed across "
                       "the predictions? The diagonal equals the recall."),
    ("Classification report", "Precision, recall, F1, support and prediction count per class plus macro "
                              "average, computed directly from the matrix."),
    ("Overall metrics", "N, accuracy, adjacent accuracy, MAE, bias, % over-/underestimated, computed from the "
                        "matrix. If exactly one experiment is selected, the values match its row in the "
                        "analysis sheet."),
    (f"Aggregation with '{ALL}'", "All matching Raw rows are summed up. Models (possibly with different layers) "
                                  "are mixed, the same text counts multiple times, large experiments dominate. "
                                  "N = number of predictions, not number of distinct texts."),
    ("Training language = single language", "Also includes the baseline of that language (there, training "
                                            "language = test language). Cross-lingual tests only: test "
                                            f"language = '{CROSS_OPTION}'."),
]

HINT_INFO = [
    ("Test set sizes", "Tests on a language that was part of training (baseline, e.g. en+de → de) use the "
                       "test split. Cross-lingual tests (e.g. de → ar) use the full dataset of the test "
                       "language, without leakage, since that language was never seen in training. Accuracy "
                       "is comparable, but cross-lingual values are statistically more stable due to larger N."),
    ("Layer per model", "If the models are evaluated on different layers, a model comparison is always also a "
                        "layer comparison."),
    ("Ordinal metrics", "Accuracy treats every confusion equally. Adjacent accuracy and MAE take into account "
                        "that B1↔B2 is less severe than A1↔C2. The bias additionally shows the direction of errors."),
    ("Bias vs. MAE", "A bias close to 0 does not mean the model is good: over- and underestimations can "
                     "cancel out. The combination of MAE (how far off) and bias (in which direction) is what matters."),
    ("Variant comparisons", "When comparing two variants (e.g. full vs. half), both should be evaluated on the "
                            "same test data. Check: test samples identical per model and test language."),
    ("Duplicate rows", "Check baseline CSVs with dedup_csvs.py (dry run) before the analysis. A text that "
                       "genuinely occurs twice in the dataset also appears twice in Raw, which is correct."),
]

UPDATE_INFO = [
    ("Change a value in Raw", "All formula sheets recalculate automatically. Raw_Paired is static and only "
                              "updates when the script is rerun."),
    ("Change a dropdown", "The confusion matrix updates immediately."),
    ("New data / new experiment", "Place the CSVs in the folder structure and rerun the script. The rows of the "
                                  "analysis sheets and the dropdown values are created by the script. Check "
                                  "the log: no unexpected [SKIP], plausible total."),
    ("Folder structure (standard)", "baseline/{test_lang}/{model}/{model}_…_l{N}.csv  and  "
                                    "crosslingual/{train_lang(s)}/[full/]testOn-{test_lang}/{model}/{model}_…_l{N}.csv. "
                                    "Other intermediate variant folders (e.g. half/) are skipped."),
    ("Folder structure (--generic)", "{variant}/…/{file}.csv. Model, training language, test language and "
                                     "layer come from the CSV, the variant from the first subfolder."),
    ("Usage", "python3 cefr_to_excel.py [FOLDER] [OUTPUT.xlsx] [--generic]"),
    ("Overwriting", "The output file is completely recreated on every run. Build your own pivots in a "
                    "separate file, otherwise they are lost on the next run."),
    ("Sorting", "Filtering in the analysis sheets is fine. Sorting via the filter breaks the block structure "
                "(blank rows / thick lines). In that case simply rerun the script."),
    ("Requirements", "Excel 2019 / Microsoft 365 or LibreOffice (because of MINIFS/MAXIFS). "
                     "No macros, no VBA."),
]


def build_info(ws, raw, has_paired):
    widths = {1: 32, 2: 30, 3: 100}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    ws["A1"] = "CEFR Results – Info & Documentation"
    ws["A1"].font = Font(bold=True, size=16)
    mode = "generic (metadata from CSV, variant = subfolder)" if GENERIC else "standard (metadata from path)"
    ws["A2"] = (f"Created on {datetime.now():%Y-%m-%d %H:%M} from '{BASE_DIR}/' – "
                f"{raw['Source_File'].nunique()} CSV files, {len(raw)} rows – mode: {mode}")
    ws["A2"].font = GRAY
    ws["A3"] = "Rows in Raw (live):"
    ws["A3"].font = BOLD
    ws["B3"] = f"=COUNTA({SH_RAW}!A:A)-1"
    ws["B3"].alignment = Alignment(horizontal="left")

    r = 3

    def height(pairs):
        lines = 1
        for text, width in pairs:
            est = sum(max(1, math.ceil(len(part) / width)) for part in str(text).split("\n"))
            lines = max(lines, est)
        return 15 * lines

    def section(title):
        nonlocal r
        r += 2
        for c in (1, 2, 3):
            ws.cell(r, c).fill = HEAD_FILL
        ws.cell(r, 1, title).font = Font(bold=True, color="FFFFFF", size=12)

    def table(header, rows):
        nonlocal r
        if header:
            r += 1
            for ci, h in enumerate(header, start=1):
                cell = ws.cell(r, ci, h)
                cell.fill, cell.font = GROUP_FILL, BOLD
        for row in rows:
            r += 1
            if len(row) == 2:   # 2 columns -> text spans B:C
                ws.cell(r, 1, row[0])
                ws.cell(r, 2, row[1])
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
                pairs = [(row[0], widths[1]), (row[1], widths[2] + widths[3])]
            else:
                for ci, v in enumerate(row, start=1):
                    ws.cell(r, ci, v)
                pairs = [(v, widths[ci]) for ci, v in enumerate(row, start=1)]
            for ci in (1, 2, 3):
                ws.cell(r, ci).alignment = WRAP
            ws.cell(r, 1).font = BOLD
            ws.row_dimensions[r].height = height(pairs)

    overview = [
        ("Core principle", f"Only the sheets '{SH_RAW}' and '{SH_PAIRED}' contain values. All other sheets "
                           "consist of Excel formulas (COUNTIFS, AVERAGEIFS, MINIFS, MAXIFS) that are "
                           "computed live from Raw."),
        ("Data source", f"All CSVs below {BASE_DIR}/. In standard mode, model, languages and layer are read "
                        "from folder and file names; if they differ from the CSV, the path wins."),
        ("Generation", "Python script (pandas + openpyxl)."),
    ]
    if GENERIC:
        overview.append(("Generic mode", "This file was created with --generic: metadata comes from the CSV, "
                                         "the first subfolder is the variant (column Training_Scope, appended "
                                         "to Training_Language). This allows direct comparison of variants "
                                         "with the same training language."))

    sheets = [s for s in SHEETS_INFO if has_paired or s[0] != SH_PAIRED]
    n = 0

    def nxt(title):
        nonlocal n
        n += 1
        section(f"{n}. {title}")

    nxt("Overview")
    table(None, overview)

    nxt("Sheets")
    table(("Sheet", "Content", "Description"), sheets)

    nxt(f"Columns in sheet '{SH_RAW}'")
    table(("Column", "Source", "Meaning"),
          [(f"{COL[c]}  –  {c}", *RAW_INFO.get(c, ("", ""))) for c in RAW_COLS])
    table(None, [
        ("Source legend", "Path = from the folder structure (in generic mode: from the CSV). File name = from "
                          "the file name. CSV = taken over unchanged. Recomputed = computed by the script from "
                          "Level_True/Level_Predicted; differences to the CSV are reported as [WARN] in the log."),
    ])

    if has_paired:
        _, _, P1, P2 = model_labels(raw)
        nxt(f"Columns in sheet '{SH_PAIRED}'")
        table(None, paired_info(P1, P2))

    nxt("Metrics (Mono_Analysis, Multi_Analysis, Confusion_Matrix)")
    table(("Metric", "Calculation", "Interpretation"), METRIC_INFO)

    nxt("Confusion matrix – usage")
    table(None, CM_INFO)

    nxt("Notes on interpretation")
    table(None, HINT_INFO)

    nxt("Updating & technical details")
    table(None, UPDATE_INFO)

    ws.sheet_view.showGridLines = False


# --------------------------------------------------------------------------
def main():
    print(f"Input: {BASE_DIR.resolve()}  |  Output: {OUT_FILE}  |  "
          f"Mode: {'generic' if GENERIC else 'standard'}\n")
    raw = load_all()

    wb = Workbook()
    ws_raw = wb.active
    ws_raw.title = SH_RAW
    write_raw(ws_raw, raw)

    # Raw_Paired directly after Raw (only with exactly two models)
    has_paired = raw["LLM_Model"].nunique() == 2
    if has_paired:
        build_paired(wb.create_sheet(SH_PAIRED), raw)
    else:
        print(f"[INFO] {SH_PAIRED} skipped (needs exactly 2 models, found: "
              f"{sorted(raw['LLM_Model'].unique())})")

    # Skip empty analysis sheets
    for ttype in ("Mono", "Multi"):
        if (raw["Training_Type"] == ttype).any():
            build_analysis(wb.create_sheet(f"{ttype}_Analysis"), raw, ttype)

    ws_cm = wb.create_sheet(SH_CM)
    refs = write_lists(wb.create_sheet(SH_LISTS), raw)
    build_confusion(ws_cm, refs)

    # Info as first sheet, active on opening
    build_info(wb.create_sheet(SH_INFO, 0), raw, has_paired)
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = (ws.title == SH_INFO)

    wb.calculation = CalcProperties(fullCalcOnLoad=True)  # Excel recalculates everything on opening
    wb.save(OUT_FILE)
    print(f"Done: {OUT_FILE.resolve()}")


if __name__ == "__main__":
    main()