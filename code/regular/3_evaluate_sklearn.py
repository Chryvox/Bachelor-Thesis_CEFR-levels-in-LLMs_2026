"""
Script 3: Evaluation (with sklearn)
===================================
Loads a trained MLP probe model (sklearn) and evaluates it on a
separate states file (different split or different language).

Usage:
  python evaluate_sklearn.py --probe probes/probe.pt --states states/test.pt
  python evaluate_sklearn.py --probe probes/probe.pt --states states/test.pt --out results/eval.xlsx
  python evaluate_sklearn.py --probe probes/probe.pt --states states/test.pt --csv results/all_results.csv
"""

import argparse
import csv
import math
import os
import numpy as np
import torch
from sklearn.metrics import classification_report, accuracy_score, confusion_matrix
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ─────────────────────────────────────────────
# CEFR mapping (fixed order, independent of label2id)
# ─────────────────────────────────────────────

CEFR_TO_NUM = {"A1": 1, "A2": 2, "B1": 3, "B2": 4, "C1": 5, "C2": 6}
_warned_labels = set()


def cefr_to_num(label, fallback_idx):
    """CEFR label -> 1..6. Falls back to index+1 if label is unknown."""
    key = str(label).strip().upper()
    if key in CEFR_TO_NUM:
        return CEFR_TO_NUM[key]
    if key not in _warned_labels:
        print(f"WARNING: Unknown CEFR label '{label}' – using index+1 as the numeric value.")
        _warned_labels.add(key)
    return int(fallback_idx) + 1


def round_half_up(x):
    """Round half up (0.5 always rounds up)."""
    return int(math.floor(float(x) + 0.5))


# ─────────────────────────────────────────────
# Load probe
# ─────────────────────────────────────────────

def load_probe(probe_path):
    probe_data = torch.load(probe_path, weights_only=False)
    model      = probe_data["probe_state"]
    label2id   = probe_data["label2id"]
    id2label   = probe_data["id2label"]

    print(f"Model loaded: {model}")
    print(f"Trained on: layer={probe_data.get('layer')}, "
          f"pooling={probe_data.get('pooling')}, "
          f"languages={probe_data.get('trained_on_languages')}")

    return model, probe_data


# ─────────────────────────────────────────────
# Additional metrics
# ─────────────────────────────────────────────

def compute_tp_fp_fn_tn(y_true, y_pred, n_classes):
    results = {}
    for c in range(n_classes):
        tp = int(np.sum((y_true == c) & (y_pred == c)))
        fp = int(np.sum((y_true != c) & (y_pred == c)))
        fn = int(np.sum((y_true == c) & (y_pred != c)))
        tn = int(np.sum((y_true != c) & (y_pred != c)))
        results[c] = {"TP": tp, "FP": fp, "FN": fn, "TN": tn}
    return results


def compute_adjacent_accuracy(y_true, y_pred, tolerance=1):
    return np.mean(np.abs(y_true - y_pred) <= tolerance)


def compute_mae(y_true, y_pred):
    return np.mean(np.abs(y_true - y_pred))


# ─────────────────────────────────────────────
# Save Excel
# ─────────────────────────────────────────────

HEADER_FILL  = PatternFill(start_color="2F4F4F", end_color="2F4F4F", fill_type="solid")
HEADER_FONT  = Font(color="FFFFFF", bold=True)
GRUEN_FILL   = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
GRUEN_FONT   = Font(bold=True)
THIN_BORDER  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin")
)

def style_header(cell):
    cell.fill      = HEADER_FILL
    cell.font      = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border    = THIN_BORDER

def style_cell(cell, bold=False):
    cell.alignment = Alignment(horizontal="center")
    cell.border    = THIN_BORDER
    if bold:
        cell.font = Font(bold=True)

def auto_width(ws):
    for col in ws.columns:
        max_len    = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def save_excel(out_path, acc, adj_acc, mae, confidence,
               report_dict, cm, id2label, label2id,
               tp_fp_etc, errors_df, probe_data, languages):

    target_names = [id2label[i] for i in sorted(id2label.keys())]
    n_total      = sum(report_dict[l]["support"] for l in target_names)

    # ── Sheet 1: Ergebnisse + TP/FP/FN/TN ──
    rows_ergebnisse = []

    # Section 1: Overview
    rows_ergebnisse.append({"Metrik": "--- Übersicht --", "Wert": ""})
    rows_ergebnisse.append({"Metrik": "Trainiert auf Sprache(n)",
                             "Wert": ", ".join(probe_data.get("trained_on_languages", []))})
    rows_ergebnisse.append({"Metrik": "Getestet auf Sprache(n)",
                             "Wert": ", ".join(sorted(set(languages)))})
    rows_ergebnisse.append({"Metrik": "Layer",       "Wert": probe_data.get("layer")})
    rows_ergebnisse.append({"Metrik": "Pooling",     "Wert": probe_data.get("pooling")})
    rows_ergebnisse.append({"Metrik": "Test-Samples","Wert": n_total})
    rows_ergebnisse.append({"Metrik": "", "Wert": ""})

    # Section 2: Metrics
    rows_ergebnisse.append({"Metrik": "--- Metriken ---", "Wert": ""})
    rows_ergebnisse.append({"Metrik": "Accuracy (Exact)",
                             "Wert": f"{acc:.4f}  ({int(acc * n_total)}/{n_total})"})
    rows_ergebnisse.append({"Metrik": "Adjacent Accuracy (±1 Level)",
                             "Wert": f"{adj_acc:.4f}  ({int(adj_acc * n_total)}/{n_total})"})
    rows_ergebnisse.append({"Metrik": "Mean Absolute Error (Level)",
                             "Wert": f"{mae:.4f}"})
    rows_ergebnisse.append({"Metrik": "Confidence (Avg)",
                             "Wert": f"{np.mean(confidence):.4f}"})
    rows_ergebnisse.append({"Metrik": "Confidence (Min)",
                             "Wert": f"{np.min(confidence):.4f}"})
    rows_ergebnisse.append({"Metrik": "Confidence (Max)",
                             "Wert": f"{np.max(confidence):.4f}"})
    rows_ergebnisse.append({"Metrik": "", "Wert": ""})

    # Section 3: Classification report
    rows_ergebnisse.append({"Metrik": "--- Klassifikationsreport ---", "Wert": ""})
    rows_ergebnisse.append({"Metrik": "Klasse",
                             "Wert": "",
                             "Precision": "Precision",
                             "Recall":    "Recall",
                             "F1-Score":  "F1-Score",
                             "Support":   "Support"})
    for label in target_names:
        r = report_dict[label]
        rows_ergebnisse.append({
            "Metrik":    label,
            "Wert":      "",
            "Precision": round(r["precision"], 4),
            "Recall":    round(r["recall"],    4),
            "F1-Score":  round(r["f1-score"],  4),
            "Support":   int(r["support"]),
        })
    # Macro avg
    r = report_dict.get("macro avg", {})
    rows_ergebnisse.append({
        "Metrik":    "Macro Avg",
        "Wert":      "",
        "Precision": round(r.get("precision", 0), 4),
        "Recall":    round(r.get("recall",    0), 4),
        "F1-Score":  round(r.get("f1-score",  0), 4),
        "Support":   int(r.get("support",     0)),
    })
    rows_ergebnisse.append({"Metrik": "", "Wert": ""})

    # Section 4: TP/FP/FN/TN
    rows_ergebnisse.append({"Metrik": "--- TP / FP / FN / TN pro Klasse ---", "Wert": ""})
    rows_ergebnisse.append({
        "Metrik": "Klasse", "Wert": "",
        "Precision": "TP", "Recall": "FP",
        "F1-Score":  "FN", "Support": "TN"
    })
    for c in sorted(id2label.keys()):
        v = tp_fp_etc[c]
        rows_ergebnisse.append({
            "Metrik":    id2label[c],
            "Wert":      "",
            "Precision": v["TP"],
            "Recall":    v["FP"],
            "F1-Score":  v["FN"],
            "Support":   v["TN"],
        })

    df_ergebnisse = pd.DataFrame(rows_ergebnisse)

    # ── Sheet 2: Konfusionsmatrix ──
    df_cm = pd.DataFrame(
        cm,
        index   =[f"Wahr: {l}" for l in target_names],
        columns =[f"Pred: {l}" for l in target_names],
    )

    # ── Sheet 3: Fehleranalyse ──
    df_errors = errors_df

    # ── Write Excel ──
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df_ergebnisse.to_excel(writer, sheet_name="Ergebnisse",       index=False)
        df_cm.to_excel(        writer, sheet_name="Konfusionsmatrix",  index=True)
        df_errors.to_excel(    writer, sheet_name="Fehleranalyse",     index=False)

    # ── Post-formatting ──
    wb = load_workbook(out_path)

    # --- Sheet 1: Ergebnisse ---
    ws = wb["Ergebnisse"]
    auto_width(ws)
    for row in ws.iter_rows():
        for cell in row:
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(horizontal="left")
            val = str(cell.value or "")
            if val.startswith("---"):
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT

    # --- Sheet 2: Konfusionsmatrix ---
    ws_cm = wb["Konfusionsmatrix"]

    # Square cells (width ≈ height)
    CELL_WIDTH  = 8
    CELL_HEIGHT = 45

    for col in range(1, len(target_names) + 2):
        ws_cm.column_dimensions[get_column_letter(col)].width = CELL_WIDTH

    for row in range(1, len(target_names) + 2):
        ws_cm.row_dimensions[row].height = CELL_HEIGHT

    for cell in ws_cm[1]:
        style_header(cell)
    for row in ws_cm.iter_rows(min_row=2):
        style_header(row[0])

    # Manual heatmap + green diagonal
    max_val = cm.max()
    n       = len(target_names)

    for i in range(n):
        for j in range(n):
            cell       = ws_cm.cell(row=i + 2, column=j + 2)
            val        = int(cm[i][j])
            cell.value     = val
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = THIN_BORDER

            if i == j:
                cell.fill = GRUEN_FILL
                cell.font = GRUEN_FONT
            else:
                intensity = val / max_val if max_val > 0 else 0
                r         = 255
                g         = int(255 - intensity * 210)
                b         = int(255 - intensity * 210)
                hex_color = f"{r:02X}{g:02X}{b:02X}"
                cell.fill = PatternFill(
                    start_color=hex_color,
                    end_color=hex_color,
                    fill_type="solid"
                )

    # --- Sheet 3: Fehleranalyse ---
    ws_err = wb["Fehleranalyse"]
    auto_width(ws_err)
    for cell in ws_err[1]:
        style_header(cell)
    for row in ws_err.iter_rows(min_row=2):
        for cell in row:
            style_cell(cell)

    wb.save(out_path)
    print(f"\n✅ Excel saved: {out_path}")


# ─────────────────────────────────────────────
# Save CSV (all texts, not just errors)
# ─────────────────────────────────────────────

CSV_COLUMNS = [
    "LLM_Model", "Layer", "Pooling",
    "Training_Language", "Test_Language",
    "Level_True", "Level_True_Num",
    "Level_Predicted", "Level_Predicted_Num",
    "Level_Distance_Abs", "Level_Distance_Signed",
    "Confidence", "Confidence_Level", "Correct",
    "Text", "states_filename", "probes_filename",
]


def save_csv(csv_path, texts, y_true, y_pred, confidence,
             languages, id2label, probe_data, test_data,
             states_path, probe_path, append=True):

    training_lang = "+".join(probe_data.get("trained_on_languages", []))
    llm_model     = test_data.get("model_id", "unknown")
    layer         = test_data.get("layer",   probe_data.get("layer"))
    pooling       = test_data.get("pooling", probe_data.get("pooling"))
    states_fname  = os.path.basename(states_path)
    probes_fname  = os.path.basename(probe_path)

    # Create target folder if needed
    out_dir = os.path.dirname(csv_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    rows = []
    for i in range(len(y_true)):
        true_idx   = int(y_true[i])
        pred_idx   = int(y_pred[i])
        true_label = id2label[true_idx]
        pred_label = id2label[pred_idx]

        true_num = cefr_to_num(true_label, true_idx)
        pred_num = cefr_to_num(pred_label, pred_idx)

        # Signed: positive = predicted too high, negative = predicted too low
        dist_signed = pred_num - true_num
        dist_abs    = abs(dist_signed)

        conf = float(confidence[i])

        rows.append({
            "LLM_Model":             llm_model,
            "Layer":                 layer,
            "Pooling":               pooling,
            "Training_Language":     training_lang,
            "Test_Language":         languages[i],
            "Level_True":            true_label,
            "Level_True_Num":        true_num,
            "Level_Predicted":       pred_label,
            "Level_Predicted_Num":   pred_num,
            "Level_Distance_Abs":    dist_abs,
            "Level_Distance_Signed": dist_signed,
            "Confidence":            round(conf, 4),
            "Confidence_Level":      round_half_up(conf * 10),
            "Correct":               int(true_num == pred_num),
            "Text":                  texts[i],
            "states_filename":       states_fname,
            "probes_filename":       probes_fname,
        })

    df = pd.DataFrame(rows, columns=CSV_COLUMNS)

    # Append mode: only write header if file doesn't exist yet
    write_header = True
    mode         = "w"
    if append and os.path.exists(csv_path):
        write_header = False
        mode         = "a"

    df.to_csv(
        csv_path,
        mode=mode,
        header=write_header,
        index=False,
        quoting=csv.QUOTE_ALL,
        encoding="utf-8-sig",   # utf-8-sig → umlauts/Arabic/Hindi open correctly in Excel
    )
    print(f"✅ CSV saved ({len(rows)} rows, mode={mode}): {csv_path}")


# ─────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--probe",  required=True, help="Path to the trained probe model")
    parser.add_argument("--states", required=True, help="Path to the test states file")
    parser.add_argument("--out",    default=None,  help="Optional path for Excel output (.xlsx)")
    parser.add_argument("--csv",    default=None,  help="Optional path for CSV output (all texts). Appended if it already exists.")
    args = parser.parse_args()

    # ── Load probe ──
    print(f"Loading probe: {args.probe}")
    probe, probe_data = load_probe(args.probe)
    label2id = probe_data["label2id"]
    id2label = probe_data["id2label"]

    # ── Load test data ──
    print(f"\nLoading test states: {args.states}")
    test_data   = torch.load(args.states, weights_only=False)
    X_test      = test_data["states"]
    cefr_levels = test_data["cefr_levels"]
    languages   = test_data["languages"]

    if X_test.dim() != 2:
        raise ValueError(f"Expected [N, hidden_dim], got: {tuple(X_test.shape)}")
    if X_test.shape[1] != probe_data["input_dim"]:
        raise ValueError(
            f"Dimension mismatch: probe expects {probe_data['input_dim']}, "
            f"states have {X_test.shape[1]}."
        )

    unknown = set(cefr_levels) - set(label2id.keys())
    if unknown:
        print(f"WARNING: Unknown labels will be ignored: {unknown}")

    keep_idx = [i for i, lvl in enumerate(cefr_levels) if lvl in label2id]
    X_test   = X_test[keep_idx].numpy()
    y_true   = np.array([label2id[cefr_levels[i]] for i in keep_idx])
    texts_kept     = [test_data["texts"][i] for i in keep_idx]
    languages_kept = [languages[i]          for i in keep_idx]

    print(f"Test set: {len(y_true)} examples, languages={sorted(set(languages))}")

    # ── Prediction ──
    print("\n--- Prediction ---")
    y_pred     = probe.predict(X_test)
    y_proba    = probe.predict_proba(X_test)
    confidence = np.max(y_proba, axis=1)

    # ── Standard metrics ──
    acc          = accuracy_score(y_true, y_pred)
    target_names = [id2label[i] for i in sorted(id2label.keys())]
    report       = classification_report(
        y_true, y_pred, target_names=target_names,
        zero_division=0, output_dict=False
    )
    report_dict  = classification_report(
        y_true, y_pred, target_names=target_names,
        zero_division=0, output_dict=True
    )
    cm = confusion_matrix(y_true, y_pred)

    # ── Ordinal metrics ──
    adj_acc = compute_adjacent_accuracy(y_true, y_pred)
    mae     = compute_mae(y_true, y_pred)

    # ── TP/FP/FN/TN ──
    tp_fp_etc = compute_tp_fp_fn_tn(y_true, y_pred, len(id2label))

    # ── Print ──
    print(f"\n=== Results ===")
    print(f"Accuracy (Exact):             {acc:.4f}  ({int(acc * len(y_true))}/{len(y_true)})")
    print(f"Adjacent Accuracy (±1 Level): {adj_acc:.4f}  ({int(adj_acc * len(y_true))}/{len(y_true)})")
    print(f"Mean Absolute Error (Level):  {mae:.4f}")
    print(f"Confidence (Avg):             {np.mean(confidence):.4f}")
    print(f"Confidence (Min):             {np.min(confidence):.4f}")
    print(f"Confidence (Max):             {np.max(confidence):.4f}")
    print("\n" + report)

    print("Confusion Matrix:")
    print(cm)

    print("\n--- TP / FP / FN / TN per class ---")
    print(f"{'Klasse':<8} {'TP':>6} {'FP':>6} {'FN':>6} {'TN':>6}")
    print("-" * 36)
    for c in sorted(id2label.keys()):
        v = tp_fp_etc[c]
        print(f"{id2label[c]:<8} {v['TP']:>6} {v['FP']:>6} {v['FN']:>6} {v['TN']:>6}")

    # ── Error analysis ──
    print("\n--- Error Analysis ---")
    errors   = y_true != y_pred
    n_errors = int(np.sum(errors))
    errors_rows = []

    if n_errors > 0:
        y_true_err = y_true[errors]
        y_pred_err = y_pred[errors]
        conf_err   = confidence[errors]
        dist_err   = np.abs(y_true_err - y_pred_err)

        n_adjacent     = int(np.sum(dist_err == 1))
        n_non_adjacent = int(np.sum(dist_err  > 1))

        print(f"Errors: {n_errors}/{len(y_true)} ({100 * n_errors / len(y_true):.2f}%)")
        print(f"  Of which adjacent  (±1):     {n_adjacent} ({100 * n_adjacent / n_errors:.1f}%)")
        print(f"  Of which non-adjacent:       {n_non_adjacent} ({100 * n_non_adjacent / n_errors:.1f}%)")
        print()

        for i, (true_idx, pred_idx) in enumerate(zip(y_true_err, y_pred_err)):
            true_label = id2label[true_idx]
            pred_label = id2label[pred_idx]
            conf       = float(conf_err[i])
            dist       = int(dist_err[i])
            benachbart = "Ja" if dist == 1 else "Nein"
            print(f"  {true_label} -> {pred_label}  (Conf: {conf:.2f}, Distance: {dist} level)")
            errors_rows.append({
                "Wahres Label":       true_label,
                "Vorhergesagtes Label": pred_label,
                "Abstand (Level)":    dist,
                "Benachbart (±1)":    benachbart,
                "Konfidenz":          round(conf, 4),
            })
    else:
        print("No errors found!")

    errors_df = pd.DataFrame(errors_rows)

    # ── Save Excel (optional) ──
    if args.out:
        save_excel(
            out_path    = args.out,
            acc         = acc,
            adj_acc     = adj_acc,
            mae         = mae,
            confidence  = confidence,
            report_dict = report_dict,
            cm          = cm,
            id2label    = id2label,
            label2id    = label2id,
            tp_fp_etc   = tp_fp_etc,
            errors_df   = errors_df,
            probe_data  = probe_data,
            languages   = languages,
        )

    # ── Save CSV (optional, all texts) ──
    if args.csv:
        save_csv(
            csv_path    = args.csv,
            texts       = texts_kept,
            y_true      = y_true,
            y_pred      = y_pred,
            confidence  = confidence,
            languages   = languages_kept,
            id2label    = id2label,
            probe_data  = probe_data,
            test_data   = test_data,
            states_path = args.states,
            probe_path  = args.probe,
            append      = True,
        )


if __name__ == "__main__":
    main()