"""
Evaluating grid search results (with normalized DE+EN ranking)
====================================================================
Reads all JSON files from the grid search folders and
creates an Excel file with all results.

Output:
    gridSearch_results.xlsx
"""

import json
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

# ─────────────────────────────────────────────
# Configuration
# ─────────────────────────────────────────────

RESULT_DIRS = {
    "DE": "/gridsearch/results/results_DE-8020",
    "EN": "/gridsearch/results/results_EN-8020",
}

OUTPUT_FILE = "gridSearch_results.xlsx"

CONFIG_COLS = ["Model", "Layer", "Architecture", "Learning Rate", "Alpha"]

# ─────────────────────────────────────────────
# Load JSON
# ─────────────────────────────────────────────

def load_json_results(result_dirs):
    """
    Loads all JSON files from the given folders.
    Returns a list of dictionaries.
    """
    all_results = []

    for sprache, ordner in result_dirs.items():
        ordner_path = Path(ordner)
        if not ordner_path.exists():
            print(f"⚠️  Folder not found: {ordner}")
            continue

        json_files = list(ordner_path.glob("*.json"))
        if not json_files:
            print(f"⚠️  No JSON files in: {ordner}")
            continue

        print(f"📂 {sprache}: {len(json_files)} JSON files found")

        for json_file in sorted(json_files):
            parts = json_file.stem.split("_")
            modell = parts[0].upper()
            layer  = parts[1]

            with open(json_file) as f:
                results = json.load(f)

            for r in results:
                all_results.append({
                    "Language":       sprache,
                    "Model":          modell,
                    "Layer":          layer,
                    "Architecture":   " -> ".join(str(l) for l in r["config"]["layers"]),
                    "Learning Rate":  r["config"]["lr"],
                    "Epochs":         r["config"]["epochs"],
                    "Alpha":          r["config"]["alpha"],
                    "Test Acc":       round(r["test_acc"], 4),
                    "Local Rank":     r.get("rank", "-"),
                })

    return all_results


# ─────────────────────────────────────────────
# Overview sheet
# ─────────────────────────────────────────────

def erstelle_uebersicht(df):
    """
    Creates a compact overview table:
    Best accuracy per language, model, and layer.
    """
    return (
        df.groupby(["Language", "Model", "Layer"])["Test Acc"]
        .max()
        .reset_index()
        .rename(columns={"Test Acc": "Best Accuracy"})
        .sort_values(["Language", "Model", "Best Accuracy"],
                     ascending=[True, True, False])
    )


# ─────────────────────────────────────────────
# Combined DE+EN ranking
# ─────────────────────────────────────────────

def erstelle_kombinierten_rank(df):
    """
    For each configuration (excluding epochs):
    - Best accuracy for DE and EN
    - Normalized score:
      Score = (DE_Acc / DE_max) + (EN_Acc / EN_max)
    - The higher the score, the better for both languages simultaneously
    - Max possible score = 2.0 (perfect for both languages)
    """
    df = df.copy()

    # Maximum accuracy per language (normalization basis)
    de_max = df[df["Language"] == "DE"]["Test Acc"].max()
    en_max = df[df["Language"] == "EN"]["Test Acc"].max()
    print(f"📊 Normalization basis: DE_max={de_max:.4f}, EN_max={en_max:.4f}")

    # Best accuracy per configuration & language (excluding epochs)
    df_de = (
        df[df["Language"] == "DE"]
        .groupby(CONFIG_COLS)
        .agg(DE_Acc=("Test Acc", "max"))
        .reset_index()
    )

    df_en = (
        df[df["Language"] == "EN"]
        .groupby(CONFIG_COLS)
        .agg(EN_Acc=("Test Acc", "max"))
        .reset_index()
    )

    # Merge on configuration
    df_merged = pd.merge(df_de, df_en, on=CONFIG_COLS, how="inner")

    # Normalized score
    df_merged["DE Norm"] = (df_merged["DE_Acc"] / de_max).round(4)
    df_merged["EN Norm"] = (df_merged["EN_Acc"] / en_max).round(4)
    df_merged["Normalized Score"] = (
        df_merged["DE Norm"] + df_merged["EN Norm"]
    ).round(4)

    # Sorting: highest score = rank 1
    df_merged = df_merged.sort_values(
        "Normalized Score", ascending=False
    ).reset_index(drop=True)
    df_merged.insert(0, "Combined Rank", df_merged.index + 1)

    final_cols = [
        "Combined Rank", "Model", "Layer", "Architecture",
        "Learning Rate", "Alpha",
        "DE_Acc", "DE Norm",
        "EN_Acc", "EN Norm",
        "Normalized Score"
    ]
    df_final = df_merged[final_cols].rename(columns={
        "DE_Acc": "DE Acc",
        "EN_Acc": "EN Acc",
    })

    # Separate tables for AYA and QWEN
    df_aya = (
        df_final[df_final["Model"] == "AYA"]
        .copy()
        .reset_index(drop=True)
    )
    df_aya["Combined Rank"] = df_aya.index + 1

    df_qwen = (
        df_final[df_final["Model"] == "QWEN"]
        .copy()
        .reset_index(drop=True)
    )
    df_qwen["Combined Rank"] = df_qwen.index + 1

    return df_final, df_aya, df_qwen


# ─────────────────────────────────────────────
# Formatting
# ─────────────────────────────────────────────

GRUEN_FILL = PatternFill(
    start_color="92D050", end_color="92D050", fill_type="solid"
)
GRUEN_FONT = Font(bold=True)


def auto_spaltenbreite(ws):
    """Automatically adjusts all column widths to fit the content."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def beste_zeile_gruen(ws, acc_col_name="Test Acc"):
    """Highlights the row with the highest accuracy in green."""
    acc_col = None
    for cell in ws[1]:
        if cell.value == acc_col_name:
            acc_col = cell.column
            break
    if acc_col is None:
        return

    best_acc = -1
    best_row = None
    for row in ws.iter_rows(min_row=2):
        val = row[acc_col - 1].value
        if val is not None and isinstance(val, (int, float)):
            if val > best_acc:
                best_acc = val
                best_row = row

    if best_row:
        for cell in best_row:
            cell.fill = GRUEN_FILL
            cell.font = GRUEN_FONT


def beste_zeile_gruen_rang(ws):
    """Highlights the row with Combined Rank = 1 in green."""
    rang_col = None
    for cell in ws[1]:
        if cell.value == "Combined Rank":
            rang_col = cell.column
            break
    if rang_col is None:
        return

    for row in ws.iter_rows(min_row=2):
        val = row[rang_col - 1].value
        if val == 1:
            for cell in row:
                cell.fill = GRUEN_FILL
                cell.font = GRUEN_FONT
            break


def heatmap_accuracy(ws, col_name="Test Acc"):
    """Colors the accuracy column as a heatmap (red → yellow → green)."""
    acc_col = None
    for cell in ws[1]:
        if cell.value == col_name:
            acc_col = cell.column
            break
    if acc_col is None:
        return

    col_letter = get_column_letter(acc_col)
    max_row = ws.max_row
    if max_row < 3:
        return

    regel = ColorScaleRule(
        start_type="min",        start_color="FF0000",
        mid_type="percentile",   mid_value=50, mid_color="FFFF00",
        end_type="max",          end_color="00B050",
    )
    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{max_row}", regel
    )


def heatmap_normalisiert(ws):
    """Heatmap on normalized score: high score = green (good)."""
    col_idx = None
    for cell in ws[1]:
        if cell.value == "Normalized Score":
            col_idx = cell.column
            break
    if col_idx is None:
        return

    col_letter = get_column_letter(col_idx)
    max_row = ws.max_row
    if max_row < 3:
        return

    # High score = green (good)
    regel = ColorScaleRule(
        start_type="min",        start_color="FF0000",  # low = red
        mid_type="percentile",   mid_value=50, mid_color="FFFF00",
        end_type="max",          end_color="00B050",    # high = green
    )
    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{max_row}", regel
    )


def formatiere_sheet(ws):
    """Applies all formatting to a normal sheet."""
    auto_spaltenbreite(ws)
    heatmap_accuracy(ws)
    beste_zeile_gruen(ws)


def formatiere_kombinations_sheet(ws):
    """Applies all formatting to a combination sheet."""
    auto_spaltenbreite(ws)
    heatmap_normalisiert(ws)
    beste_zeile_gruen_rang(ws)


# ─────────────────────────────────────────────
# Create Excel
# ─────────────────────────────────────────────

def erstelle_excel(all_results, output_file):
    df = pd.DataFrame(all_results)

    # Calculate global rank
    df = df.sort_values("Test Acc", ascending=False).reset_index(drop=True)
    df["Global Rank"]       = df.index + 1
    df["Total"]             = len(df)
    df["Global Rank (X/Y)"] = (
        df["Global Rank"].astype(str) + "/" + df["Total"].astype(str)
    )

    uebersicht = erstelle_uebersicht(df)

    # Calculate combined ranking
    df_kombi, df_aya_kombi, df_qwen_kombi = erstelle_kombinierten_rank(df)

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:

        # ── Sheet 1: Overview ──
        uebersicht.to_excel(writer, sheet_name="Overview", index=False)
        print(f"✅ Sheet 'Overview' created ({len(uebersicht)} rows)")

        # ── Sheet 2: All results ──
        df_alle = df[[
            "Global Rank (X/Y)", "Language", "Model", "Layer",
            "Architecture", "Learning Rate", "Epochs", "Alpha",
            "Test Acc", "Local Rank"
        ]].copy()
        df_alle.to_excel(writer, sheet_name="All Results", index=False)
        print(f"✅ Sheet 'All Results' created ({len(df_alle)} rows)")

        # ── Sheet 3: Combined ranking (all models) ──
        df_kombi.to_excel(writer, sheet_name="Combined DE+EN", index=False)
        print(f"✅ Sheet 'Combined DE+EN' created ({len(df_kombi)} rows)")

        # ── Sheet 4: AYA combined ──
        df_aya_kombi.to_excel(
            writer, sheet_name="AYA Combined DE+EN", index=False
        )
        print(
            f"✅ Sheet 'AYA Combined DE+EN' created "
            f"({len(df_aya_kombi)} rows)"
        )

        # ── Sheet 5: QWEN combined ──
        df_qwen_kombi.to_excel(
            writer, sheet_name="QWEN Combined DE+EN", index=False
        )
        print(
            f"✅ Sheet 'QWEN Combined DE+EN' created "
            f"({len(df_qwen_kombi)} rows)"
        )

        # ── Sheets per language+model+layer ──
        for (sprache, modell, layer), gruppe in df.groupby(
            ["Language", "Model", "Layer"]
        ):
            sheet_name = f"{sprache}_{modell}_{layer}"
            gruppe_sorted = gruppe.sort_values("Test Acc", ascending=False)
            gruppe_sorted[[
                "Language", "Model", "Layer", "Architecture",
                "Learning Rate", "Epochs", "Alpha", "Test Acc",
                "Local Rank", "Global Rank (X/Y)"
            ]].to_excel(writer, sheet_name=sheet_name, index=False)
            print(
                f"✅ Sheet '{sheet_name}' created "
                f"({len(gruppe_sorted)} rows)"
            )

    # ── Apply formatting afterwards ──
    print("\n🎨 Applying formatting...")
    wb = load_workbook(output_file)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if "Combined" in sheet_name:
            formatiere_kombinations_sheet(ws)
        else:
            formatiere_sheet(ws)
        print(f"   ✅ {sheet_name} formatted")

    wb.save(output_file)
    print(f"\n🎉 Excel saved: {output_file}")


# ─────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────

def main():
    print("Loading grid search results...\n")
    all_results = load_json_results(RESULT_DIRS)

    if not all_results:
        print("❌ No results found!")
        return

    print(f"\n📊 Total: {len(all_results)} configurations loaded")
    print("Creating Excel...\n")
    erstelle_excel(all_results, OUTPUT_FILE)


if __name__ == "__main__":
    main()